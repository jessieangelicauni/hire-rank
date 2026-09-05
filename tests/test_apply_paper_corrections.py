import io

import docx
import pytest
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

from scripts.apply_paper_corrections import (
    ABSTRACT_TEXT,
    add_evidence_table_captions,
    apply_table_and_figure_layout_fixes,
    find_paragraph,
    fix_abstract_and_contributions,
    fix_conclusion_and_limitations,
    fix_experimental_setup_and_table_ii,
    fix_methods_iii_c_d_e,
    insert_paragraph_before,
    insert_results_headers_and_ranking_failure_table,
    insert_table_before,
    insert_weakness_retry_audit_table,
    replace_paragraph_text,
    swap_figure_2_image,
)

_OLD_ABSTRACT = (
    "Abstract—Talent-acquisition systems increasingly pair large language models (LLMs) with statistical ranking "
    "to shortlist and order job applicants. The most closely related recent system combines an LLM-driven active "
    "listwise tournament with Plackett-Luce aggregation for human-resources applicant ranking, but it reports no "
    "mechanism against LLM identifier-drift failures during ranking, gives no cross-run stability, and offers "
    "only prompt-level mitigation against hallucinated assessment claims. This paper's objective is to validate "
    "three pipeline-level mechanisms that close these gaps in an open, reproducible applicant-ranking system. "
    "This research's methods are: 1) a LangGraph-orchestrated pipeline that shortlists applicants by "
    "embedding-based skill overlap; 2) resume-grounded strength/weakness assessments generated with an automated "
    "self-correction retry loop; and 3) ranking of shortlisted applicants with a position-robust, "
    "schema-constrained listwise tournament, adapting the Monte Carlo knowledge-gradient (MC-KG) "
    "subset-selection rule and Bayesian Plackett-Luce aggregation from prior work. The results show that mean "
    "Faithfulness reached 0.896 and mean within-repeat ranking convergence (Kendall-τ) reached 0.953, both "
    "measured across all 10 job profiles. In conclusion, these results — measured with metrics the closest prior "
    "system does not report or disclose — offer transparent, empirical evidence that this research addresses "
    "those gaps."
)


def test_fix_abstract_and_contributions_removes_cross_run_stability_claim():
    document = docx.Document()
    document.add_paragraph(_OLD_ABSTRACT)

    fix_abstract_and_contributions(document)

    paragraph = document.paragraphs[0]
    assert "cross-run stability" not in paragraph.text
    assert paragraph.text == ABSTRACT_TEXT


def test_fix_abstract_and_contributions_lists_the_three_introduction_contributions_in_order():
    document = docx.Document()
    document.add_paragraph(_OLD_ABSTRACT)

    fix_abstract_and_contributions(document)

    text = document.paragraphs[0].text
    self_correcting_idx = text.index("self-correcting" if "self-correcting" in text else "self-correction")
    position_robust_idx = text.index("position-robust")
    adaptive_idx = text.index("pool-size-aware adaptive iteration")
    assert self_correcting_idx < position_robust_idx < adaptive_idx


_OLD_FIG2_CAPTION = (
    "As shown, each generated claim is checked against evidence retrieved from the resume for unsupported, "
    "contradicted, or unverifiable content, with the assessment regenerated once on any detected issue before "
    "it is accepted."
)

_OLD_ADAPTIVE_ITERATION_PARAGRAPH = (
    "The pipeline's third contribution addresses the fixed-iteration-count gap identified in Section I: at a "
    "fixed iteration count, a large shortlist receives too few comparisons per applicant to rank reliably, while "
    "a small shortlist receives far more than it needs. Rather than using a fixed count for every job profile "
    "regardless of shortlist size, the iteration count is scaled so that every shortlisted applicant is expected "
    "to appear in roughly a fixed target number of sampled subsets on average, subject to a floor and a ceiling."
)

_OLD_SECTION_III_E = (
    "Each job profile's final ranking averages applicant utilities across its successful stability repeats, "
    "which reduces the noise a single repeat's point estimate would otherwise carry into the ranking. "
    "Within-repeat convergence — via the Kendall-τ rank correlation between the utility ordering at consecutive "
    "iterations — is computed directly from the tournament traces, with no LLM call required. Separately, an "
    "offline harness audits assessment quality after the fact: strengths are scored with the Faithfulness metric "
    "(an LLM judge estimates how well each strength is entailed by the resume text, given a neutral, "
    "job-agnostic question rather than one naming the job title, so the score reflects only resume support and "
    "not job-profile-specific framing), while weaknesses — which are almost always absence claims that a "
    "faithfulness-style entailment check cannot verify positively — are instead audited with the same "
    "skill-contradiction heuristic used in the live retry loop, applied here against the job profile's required "
    "skills rather than the applicant's own skill list."
)


def test_fix_methods_rewrites_figure_2_caption():
    document = docx.Document()
    document.add_paragraph(_OLD_FIG2_CAPTION)
    document.add_paragraph(_OLD_ADAPTIVE_ITERATION_PARAGRAPH)
    document.add_paragraph(_OLD_SECTION_III_E)

    fix_methods_iii_c_d_e(document)

    caption_text = document.paragraphs[0].text
    assert "evidence retrieved from the resume" not in caption_text
    assert "each generated weakness is checked" in caption_text
    assert "Strengths and additional skills are not live-checked" in caption_text


def test_fix_methods_adds_adaptive_iteration_formula():
    document = docx.Document()
    document.add_paragraph(_OLD_FIG2_CAPTION)
    document.add_paragraph(_OLD_ADAPTIVE_ITERATION_PARAGRAPH)
    document.add_paragraph(_OLD_SECTION_III_E)

    fix_methods_iii_c_d_e(document)

    formula_text = document.paragraphs[1].text
    assert "ceil(target_appearances_per_candidate" in formula_text
    assert "124 iterations" in formula_text
    assert "30 iterations" in formula_text


def test_fix_methods_corrects_weakness_audit_ground_truth():
    document = docx.Document()
    document.add_paragraph(_OLD_FIG2_CAPTION)
    document.add_paragraph(_OLD_ADAPTIVE_ITERATION_PARAGRAPH)
    document.add_paragraph(_OLD_SECTION_III_E)

    fix_methods_iii_c_d_e(document)

    audit_text = document.paragraphs[2].text
    assert "applied here against the job profile's required skills rather than the applicant's own skill list" not in audit_text
    assert "job-profile-skill-bridging half" in audit_text
    assert "which is what actually decides whether the weakness is a genuine contradiction" in audit_text


def test_find_paragraph_returns_matching_paragraph():
    document = docx.Document()
    document.add_paragraph("Some other text.")
    target = document.add_paragraph("This contains the marker phrase.")

    found = find_paragraph(document, "marker phrase")

    # python-docx constructs a brand-new Paragraph wrapper object on every
    # access of `.paragraphs` (see docx/blkcntnr.py: `[Paragraph(p, self) for
    # p in self._element.p_lst]`), so wrapper identity (`found is target`) is
    # never preserved even when both wrap the same underlying paragraph.
    # `._p` is the underlying lxml element, whose identity IS stable, and is
    # what actually matters: mutating through `found` mutates the same node
    # `target` wraps.
    assert found._p is target._p


def test_find_paragraph_raises_when_not_found():
    document = docx.Document()
    document.add_paragraph("Nothing relevant here.")

    try:
        find_paragraph(document, "missing marker")
        assert False, "expected ValueError"
    except ValueError:
        pass


def test_replace_paragraph_text_single_run():
    document = docx.Document()
    paragraph = document.add_paragraph("Old sentence to replace.")

    replace_paragraph_text(paragraph, "New sentence.")

    assert paragraph.text == "New sentence."


def test_replace_paragraph_text_preserves_bold_formatting():
    document = docx.Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Old bold sentence.")
    run.bold = True

    replace_paragraph_text(paragraph, "New bold sentence.")

    assert paragraph.text == "New bold sentence."
    assert paragraph.runs[0].bold is True


def test_replace_paragraph_text_multi_run_collapses_to_one():
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Part one. ")
    paragraph.add_run("Part two.")

    replace_paragraph_text(paragraph, "Single replacement sentence.")

    assert paragraph.text == "Single replacement sentence."


def test_insert_paragraph_before_places_new_paragraph_immediately_before_anchor():
    document = docx.Document()
    document.add_paragraph("Before anchor.")
    anchor = document.add_paragraph("Anchor paragraph.")
    document.add_paragraph("After anchor.")

    inserted = insert_paragraph_before(anchor, "Inserted heading", bold=True)

    texts = [p.text for p in document.paragraphs]
    assert texts == ["Before anchor.", "Inserted heading", "Anchor paragraph.", "After anchor."]
    assert inserted.runs[0].bold is True


def test_insert_table_before_places_table_immediately_before_anchor():
    document = docx.Document()
    document.add_paragraph("Before anchor.")
    anchor = document.add_paragraph("Anchor paragraph.")

    table = insert_table_before(document, anchor, rows=2, cols=2, col_widths_in=[1.5, 1.5])
    table.cell(0, 0).text = "Header"

    body_children = list(document.element.body)
    tbl_index = next(i for i, el in enumerate(body_children) if el.tag.endswith("}tbl"))
    anchor_index = next(i for i, el in enumerate(body_children) if el is anchor._p)
    assert tbl_index == anchor_index - 1
    assert table.cell(0, 0).text == "Header"


def test_insert_table_before_with_col_widths_disables_autofit():
    document = docx.Document()
    anchor = document.add_paragraph("Anchor paragraph.")

    table = insert_table_before(document, anchor, rows=2, cols=2, col_widths_in=[1.5, 1.5])

    assert table.autofit is False
    tbl_layout = table._tbl.tblPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblLayout"
    )
    assert tbl_layout is not None
    assert tbl_layout.get(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type"
    ) == "fixed"


def test_insert_table_before_without_col_widths_keeps_default_autofit():
    document = docx.Document()
    anchor = document.add_paragraph("Anchor paragraph.")

    table = insert_table_before(document, anchor, rows=2, cols=2)

    assert table.autofit is True


_OLD_EXPERIMENTAL_SETUP_START = (
    "All results runs against the Qwen2.5-14B-Instruct model, over the repository's corpus of resumes (500 "
    "sampled from the EraMatch CV Parsing Benchmark v3.0 synthetic resume dataset [14]) and 10 job profiles, "
    "authored synthetically by the authors. Shortlisting (Section III-B) narrowed this corpus to 348 (job, "
    "applicant) pairs across the 10 profiles."
)

_OLD_TABLE_II_CAPTION = (
    "Table II. Faithfulness scores for generated strengths, from the full-coverage audit against the latest "
    "assessments (run-id 20260831-010721), covering all 10 job profiles (2 evaluation failures excluded from "
    "the 300-item sample). The run-wide mean is 0.896, following successive prompt refinements described in "
    "Section III-C."
)


def test_fix_experimental_setup_grammar():
    document = docx.Document()
    document.add_paragraph(_OLD_EXPERIMENTAL_SETUP_START)
    document.add_paragraph(_OLD_TABLE_II_CAPTION)

    fix_experimental_setup_and_table_ii(document)

    assert "All results runs against" not in document.paragraphs[0].text
    assert "All results were obtained by running the pipeline against" in document.paragraphs[0].text


def test_fix_table_ii_caption_consistency():
    document = docx.Document()
    document.add_paragraph(_OLD_EXPERIMENTAL_SETUP_START)
    document.add_paragraph(_OLD_TABLE_II_CAPTION)

    fix_experimental_setup_and_table_ii(document)

    caption = document.paragraphs[1].text
    assert "full-coverage audit" not in caption
    assert "300-item stratified sample" in caption
    assert "298 of the 300 sampled items scored" in caption


def _results_section_skeleton():
    document = docx.Document()
    document.add_paragraph("V. Results and Discussion")
    document.add_paragraph("This section reports convergence, per-role stability, and generation faithfulness.")
    fig4 = document.add_paragraph("Fig. 4. Within-repeat convergence (Kendall-τ) across the 10 job profiles.")
    document.add_paragraph("Fig. 5. Final Kendall's Tau per job role after convergence.")
    table_ii_caption = document.add_paragraph(
        "Table II. Faithfulness scores for generated strengths, from a 300-item stratified sample."
    )
    fig6 = document.add_paragraph("Fig. 6. Faithfulness scores for generated strengths, by job profile.")
    comparison_header = document.add_paragraph("C. Comparison with the Closest Prior System")
    # The real document's existing "A./B./C./D./E." subsection headers (e.g.
    # Section III's) are bold *and* italic -- match that here so the test can
    # actually catch a newly-inserted header that doesn't match the style.
    comparison_header.runs[0].bold = True
    comparison_header.runs[0].italic = True
    return document, fig4, table_ii_caption, comparison_header


def test_insert_results_headers_adds_a_and_b_before_the_right_content():
    document, fig4, table_ii_caption, comparison_header = _results_section_skeleton()

    insert_results_headers_and_ranking_failure_table(document)

    texts = [p.text for p in document.paragraphs]
    assert "A. Tournament Convergence" in texts
    assert texts.index("A. Tournament Convergence") == texts.index(fig4.text) - 1
    assert "B. Generation Faithfulness and Contradiction Audit" in texts
    assert texts.index("B. Generation Faithfulness and Contradiction Audit") == texts.index(table_ii_caption.text) - 1

    header_paragraphs = {p.text: p for p in document.paragraphs}
    for header_text in ("A. Tournament Convergence", "B. Generation Faithfulness and Contradiction Audit"):
        run = header_paragraphs[header_text].runs[0]
        assert run.bold == comparison_header.runs[0].bold
        assert run.italic == comparison_header.runs[0].italic


def test_insert_results_adds_ranking_failure_table_before_comparison_section():
    document, fig4, table_ii_caption, comparison_header = _results_section_skeleton()

    insert_results_headers_and_ranking_failure_table(document)

    body_children = list(document.element.body)
    tbl_indices = [i for i, el in enumerate(body_children) if el.tag.endswith("}tbl")]
    comparison_index = next(i for i, el in enumerate(body_children) if el is comparison_header._p)
    assert len(tbl_indices) == 1
    assert tbl_indices[0] < comparison_index

    table = document.tables[0]
    assert table.cell(0, 0).text == "Metric"
    assert table.cell(1, 1).text == "1713"
    assert table.cell(4, 1).text == "0"


_OLD_CONCLUSION = (
    "This paper presented three pipeline-level mechanisms for LLM-driven applicant ranking, validated against a "
    "real, locally-hosted 14-billion-parameter model rather than an undisclosed one. Self-correcting assessment "
    "generation achieves a 0.0% live skill-contradiction rate (0.6% under independent audit); the "
    "position-robust tournament completes reliably, with zero unrecovered ranking failures; and adaptive "
    "iteration scaling settles convergence at Kendall-τ 0.91–0.98 across all ten job profiles, a reproducible "
    "number the closest prior system's rescaled curve never discloses. Mean Faithfulness reaches 0.896, with "
    "the residual gap traced to Faithfulness's own architecture rather than generation quality."
)

_OLD_FUTURE_WORK = (
    "Building on these results, future work will pursue two directions: 1) human-rater validation — comparing "
    "this pipeline against independent human-expert rankings using the same protocol as [1]; and 2) an "
    "independent judge model — sourcing Faithfulness scoring from a separate model than the one used for "
    "generation, to rule out shared-model bias. Taken together, these mechanisms move an already-strong "
    "architecture — the LLM listwise tournament with Plackett-Luce aggregation — toward one whose failure modes "
    "on a real, imperfect local model are documented, fixed, and measured, with the future work above set to "
    "close the remaining validation gaps."
)


def test_fix_conclusion_renames_independent_audit_and_softens_faithfulness_claim():
    document = docx.Document()
    document.add_paragraph(_OLD_CONCLUSION)
    document.add_paragraph(_OLD_FUTURE_WORK)

    fix_conclusion_and_limitations(document)

    conclusion_text = document.paragraphs[0].text
    assert "under independent audit" not in conclusion_text
    assert "0.6%" not in conclusion_text
    assert "offline post-hoc audit" in conclusion_text
    assert "0.82%" in conclusion_text
    assert "traced to Faithfulness's own architecture rather than generation quality" not in conclusion_text
    assert "cannot rule out that part of the gap reflects genuine generation quality" in conclusion_text


def test_fix_conclusion_inserts_limitations_paragraph_and_extends_future_work():
    document = docx.Document()
    document.add_paragraph(_OLD_CONCLUSION)
    document.add_paragraph(_OLD_FUTURE_WORK)

    fix_conclusion_and_limitations(document)

    texts = [p.text for p in document.paragraphs]
    limitations = [t for t in texts if "synthetic resumes" in t]
    assert len(limitations) == 1
    assert "real-world hiring validity" in limitations[0]
    assert "no demographic or fairness analysis" in limitations[0]

    future_work_text = [t for t in texts if "future work will pursue" in t][0]
    assert "three directions" in future_work_text
    assert "demographic and fairness evaluation" in future_work_text
    limitations_index = texts.index(limitations[0])
    future_work_index = texts.index(future_work_text)
    assert limitations_index < future_work_index


def _tiny_png_path(tmp_path):
    path = tmp_path / "tiny.png"
    Image.new("RGB", (100, 40), color="white").save(path)
    return path


def _colored_png_path(tmp_path, name: str, color: str):
    path = tmp_path / name
    Image.new("RGB", (10, 10), color=color).save(path)
    return path


def _inline_shape_rel_id(shape):
    return shape._inline.graphic.graphicData.pic.blipFill.blip.embed


def test_swap_figure_2_image_replaces_only_the_matching_dimension_shape(tmp_path):
    document = docx.Document()
    document.add_picture(str(_colored_png_path(tmp_path, "first.png", "red")), width=Inches(2.0), height=Inches(1.0))
    document.add_picture(
        str(_colored_png_path(tmp_path, "second.png", "blue")), width=Inches(3.23), height=Inches(1.90)
    )
    new_image_path = _colored_png_path(tmp_path, "new.png", "green")
    backup_path = tmp_path / "archive" / "image2.original.png"

    first_rel_id = _inline_shape_rel_id(document.inline_shapes[0])
    second_rel_id = _inline_shape_rel_id(document.inline_shapes[1])
    original_first_bytes = document.part.related_parts[first_rel_id].blob
    original_second_bytes = document.part.related_parts[second_rel_id].blob

    swap_figure_2_image(document, new_image_path=new_image_path, backup_path=backup_path)

    assert document.part.related_parts[first_rel_id].blob == original_first_bytes
    assert document.part.related_parts[second_rel_id].blob == new_image_path.read_bytes()
    assert document.part.related_parts[second_rel_id].blob != original_second_bytes
    assert backup_path.exists()
    assert backup_path.read_bytes() == original_second_bytes


def test_swap_figure_2_image_raises_on_dimension_mismatch_and_does_not_modify_anything(tmp_path):
    document = docx.Document()
    document.add_picture(str(_colored_png_path(tmp_path, "first.png", "red")), width=Inches(2.0), height=Inches(1.0))
    # Second shape deliberately does NOT match the expected ~3.23in x 1.90in.
    document.add_picture(str(_colored_png_path(tmp_path, "second.png", "blue")), width=Inches(5.0), height=Inches(4.0))
    new_image_path = _colored_png_path(tmp_path, "new.png", "green")
    backup_path = tmp_path / "archive" / "image2.original.png"

    first_rel_id = _inline_shape_rel_id(document.inline_shapes[0])
    second_rel_id = _inline_shape_rel_id(document.inline_shapes[1])
    original_first_bytes = document.part.related_parts[first_rel_id].blob
    original_second_bytes = document.part.related_parts[second_rel_id].blob

    with pytest.raises(ValueError):
        swap_figure_2_image(document, new_image_path=new_image_path, backup_path=backup_path)

    assert document.part.related_parts[first_rel_id].blob == original_first_bytes
    assert document.part.related_parts[second_rel_id].blob == original_second_bytes
    assert not backup_path.exists()


def test_swap_figure_2_image_rejects_a_near_miss_within_the_old_looser_tolerance(tmp_path):
    document = docx.Document()
    document.add_picture(str(_colored_png_path(tmp_path, "first.png", "red")), width=Inches(2.0), height=Inches(1.0))
    # 0.03in off on height -- matches the real paper's Figure 3, which sits
    # this close to Figure 2's dimensions. Must be rejected: a looser
    # tolerance (e.g. 0.1in) would have accepted this as a false match.
    document.add_picture(
        str(_colored_png_path(tmp_path, "second.png", "blue")), width=Inches(3.23), height=Inches(1.87)
    )
    new_image_path = _colored_png_path(tmp_path, "new.png", "green")
    backup_path = tmp_path / "archive" / "image2.original.png"

    second_rel_id = _inline_shape_rel_id(document.inline_shapes[1])
    original_second_bytes = document.part.related_parts[second_rel_id].blob

    with pytest.raises(ValueError):
        swap_figure_2_image(document, new_image_path=new_image_path, backup_path=backup_path)

    assert document.part.related_parts[second_rel_id].blob == original_second_bytes
    assert not backup_path.exists()


def test_apply_table_and_figure_layout_fixes(tmp_path):
    document = docx.Document()
    # The real paper's body section is 2-column; a fresh python-docx
    # Document() defaults to 1 column, which would make the "widen to 1
    # column" fix a no-op and hide a swapped-marker bug entirely. Force a
    # 2-column body section here so this test actually exercises widening.
    # python-docx's default template's sectPr already has a <w:cols>
    # element (with no explicit w:num, i.e. 1 column) -- modify it in place
    # rather than appending a second one, since find() would keep returning
    # the original, unmodified element.
    existing_cols = document.sections[0]._sectPr.find(qn("w:cols"))
    if existing_cols is None:
        existing_cols = OxmlElement("w:cols")
        document.sections[0]._sectPr.append(existing_cols)
    existing_cols.set(qn("w:num"), "2")

    # Two unrelated tables inserted earlier in the document, mimicking the
    # ranking-failure and weakness-retry-audit tables Tasks 10/12 add before
    # Table III's position -- Table III must not be found by position.
    document.add_table(rows=2, cols=2)
    document.add_table(rows=2, cols=2)
    # These two paragraphs mirror the real document's structure: the
    # ranking-failure-evidence intro precedes "C. Comparison...", which in
    # turn precedes Table III -- apply_table_and_figure_layout_fixes widens
    # everything from the first anchor through (not including) the second
    # as one full-width block, then separately widens Table III.
    ranking_failure_intro = document.add_paragraph(
        "The position-robust tournament's reliability claim is backed by the following counts from the full run."
    )
    comparison_paragraph = document.add_paragraph("C. Comparison with the Closest Prior System")
    table_iii = document.add_table(rows=2, cols=3)
    table_iii.cell(0, 0).text = "Dimension"
    document.add_picture(str(_tiny_png_path(tmp_path)), width=Inches(3.4))

    apply_table_and_figure_layout_fixes(document)

    for table in document.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            assert tr_pr.find(qn("w:cantSplit")) is not None
        header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
        assert header_tr_pr.find(qn("w:tblHeader")) is not None

    assert table_iii.columns[0].width == Inches(1.6)
    assert table_iii.columns[1].width == Inches(3.0)
    assert table_iii.columns[2].width == Inches(3.0)

    def _sect_pr(element):
        sect_pr = element.find(qn("w:pPr") + "/" + qn("w:sectPr"))
        assert sect_pr is not None
        return sect_pr

    def _sect_pr_cols_num(element):
        return _sect_pr(element).find(qn("w:cols")).get(qn("w:num"))

    def _sect_pr_type(element):
        type_el = _sect_pr(element).find(qn("w:type"))
        return type_el.get(qn("w:val")) if type_el is not None else None

    body_children = list(document.element.body)
    tbl_index = next(i for i, el in enumerate(body_children) if el is table_iii._tbl)
    # The marker *before* the table ends the preceding section -- it must
    # keep the document's original column count (2), not the new width.
    assert _sect_pr_cols_num(body_children[tbl_index - 1]) == "2"
    # The marker *after* the table ends the section containing the table
    # itself -- per OOXML, a paragraph's sectPr governs backwards, so this
    # is the one that must actually be single-column for the table to render
    # at full width instead of being squeezed into one of the two columns.
    assert _sect_pr_cols_num(body_children[tbl_index + 1]) == "1"
    # Tables get a page-break (not continuous) section change around them --
    # LibreOffice's PDF export garbles/overlaps a continuous section change
    # that directly abuts a table.
    assert _sect_pr_type(body_children[tbl_index + 1]) == "nextPage"

    # The ranking-failure-evidence/weakness-retry-audit block (everything
    # from the ranking-failure intro up to "C. Comparison...") is widened as
    # a single full-width range, not per-table -- one pair of markers
    # bracketing the whole range, not one pair per table inside it.
    intro_index = next(i for i, el in enumerate(body_children) if el is ranking_failure_intro._p)
    comparison_index = next(i for i, el in enumerate(body_children) if el is comparison_paragraph._p)
    assert _sect_pr_cols_num(body_children[intro_index - 1]) == "2"
    assert _sect_pr_cols_num(body_children[comparison_index - 1]) == "1"
    assert _sect_pr_type(body_children[comparison_index - 1]) == "nextPage"

    figure_1_shape = document.inline_shapes[0]
    assert figure_1_shape.width == Inches(6.8)

    figure_1_paragraph = figure_1_shape._inline.getparent().getparent().getparent()
    fig_index = next(i for i, el in enumerate(body_children) if el is figure_1_paragraph)
    assert _sect_pr_cols_num(body_children[fig_index - 1]) == "2"
    assert _sect_pr_cols_num(body_children[fig_index + 1]) == "1"
    assert _sect_pr_type(body_children[fig_index + 1]) == "continuous"


def _table_declared_width_in(table) -> float:
    """Sums a table's declared column widths (as written to <w:tblGrid>'s
    <w:gridCol> elements via table.columns[i].width) and converts to
    inches. A table that was never given explicit widths (autofit, no
    gridCol width attribute) contributes 0 -- it can't be "wide" by
    declared value, so it's correctly excluded from the invariant below."""
    total_emu = 0
    for column in table.columns:
        if column.width is not None:
            total_emu += column.width
    return total_emu / 914400


def _num_columns_governing_index(body, body_children, index) -> int:
    """Returns the column count of the section that governs the body
    element at `index`, using the same OOXML semantics
    _widen_to_full_page_width's docstring documents and relies on: a
    paragraph's <w:pPr>/<w:sectPr> describes the section ENDING at that
    paragraph (i.e. it governs everything back to the previous section
    break, not what follows it). So the section governing a given element
    is defined by the *next* paragraph-level sectPr at or after it in body
    order, falling back to the body's own trailing sectPr if none follows
    before the end of the document."""
    for element in body_children[index:]:
        if not element.tag.endswith("}p"):
            continue
        p_pr = element.find(qn("w:pPr"))
        sect_pr = p_pr.find(qn("w:sectPr")) if p_pr is not None else None
        if sect_pr is not None:
            cols = sect_pr.find(qn("w:cols"))
            num_attr = cols.get(qn("w:num")) if cols is not None else None
            return int(num_attr) if num_attr is not None else 1
    final_sect_pr = body.find(qn("w:sectPr"))
    if final_sect_pr is not None:
        cols = final_sect_pr.find(qn("w:cols"))
        num_attr = cols.get(qn("w:num")) if cols is not None else None
        return int(num_attr) if num_attr is not None else 1
    return 1


def test_wide_tables_are_always_bracketed_by_a_widened_section(tmp_path):
    """Structural regression test for the wide-table-in-narrow-column bug
    class (Task 14 bug 4). test_apply_table_and_figure_layout_fixes above
    only checks that the two currently-hardcoded wide tables/blocks get
    widened; it would NOT catch a future task that adds a third wide table
    without adding a corresponding widening call. This test instead checks
    the general invariant apply_table_and_figure_layout_fixes is supposed
    to uphold: any table whose declared total column width exceeds one
    column's worth of the document's original multi-column body must sit
    in a section the fixes have widened to 1 column -- computed
    structurally from the actual section/table XML, not by name-checking
    the two known tables.
    """
    document = docx.Document()
    existing_cols = document.sections[0]._sectPr.find(qn("w:cols"))
    if existing_cols is None:
        existing_cols = OxmlElement("w:cols")
        document.sections[0]._sectPr.append(existing_cols)
    existing_cols.set(qn("w:num"), "2")
    original_num_cols = 2

    # Two unrelated, unwidened, genuinely-narrow tables. python-docx's
    # add_table default (no explicit widths) actually declares 3.0in per
    # column, i.e. 6.0in total for a 2-column table -- which would itself
    # trip the "wider than one narrow column" check below despite being an
    # ordinary, correctly-unwidened table -- so give them explicit narrow
    # widths to exercise the "must NOT trip the invariant" case cleanly.
    for _ in range(2):
        narrow_table = document.add_table(rows=2, cols=2)
        narrow_table.autofit = False
        for column in narrow_table.columns:
            column.width = Inches(1.0)

    # A stand-in for the ranking-failure/weakness-retry-audit evidence block:
    # a 6.0in-wide table sitting between two paragraphs that
    # apply_table_and_figure_layout_fixes widens as a single ranged block
    # (not via a marker adjacent to the table itself).
    ranking_failure_intro = document.add_paragraph(
        "The position-robust tournament's reliability claim is backed by the following counts from the full run."
    )
    comparison_paragraph = document.add_paragraph("C. Comparison with the Closest Prior System")
    insert_table_before(document, comparison_paragraph, rows=2, cols=2, col_widths_in=[4.5, 1.5])

    table_iii = document.add_table(rows=2, cols=3)
    table_iii.cell(0, 0).text = "Dimension"
    document.add_picture(str(_tiny_png_path(tmp_path)), width=Inches(3.4))

    apply_table_and_figure_layout_fixes(document)

    reference_section = document.sections[0]
    content_width_in = (
        reference_section.page_width - reference_section.left_margin - reference_section.right_margin
    ) / 914400
    narrow_column_width_in = content_width_in / original_num_cols

    body = document.element.body
    body_children = list(body)
    checked_a_wide_table = False
    for table in document.tables:
        declared_width_in = _table_declared_width_in(table)
        if declared_width_in <= narrow_column_width_in:
            continue
        checked_a_wide_table = True
        index = body_children.index(table._tbl)
        num_cols = _num_columns_governing_index(body, body_children, index)
        assert num_cols == 1, (
            f"table declared at {declared_width_in:.2f}in exceeds the original "
            f"{narrow_column_width_in:.2f}in-wide column but sits in a {num_cols}-column "
            f"section instead of a widened 1-column one"
        )

    # Sanity check the test itself actually exercised the invariant against
    # at least one wide table -- otherwise a bug that deletes every widening
    # call could pass here vacuously.
    assert checked_a_wide_table


_FAKE_REPORT = {
    "total_pairs": 348,
    "total_weaknesses_checked": 1298,
    "pairs_with_initial_contradiction": 135,
    "pairs_fixed_by_retry": 97,
    "pairs_dropped": 38,
    "items_dropped": 49,
    "offline_audit_residual_contradictions": 10,
    "offline_audit_denominator": 1224,
    "offline_audit_rate": 10 / 1224,
}


def test_insert_weakness_retry_audit_table_before_comparison_section():
    document, fig4, table_ii_caption, comparison_header = _results_section_skeleton()

    insert_weakness_retry_audit_table(document, _FAKE_REPORT)

    body_children = list(document.element.body)
    tbl_index = next(i for i, el in enumerate(body_children) if el.tag.endswith("}tbl"))
    comparison_index = next(i for i, el in enumerate(body_children) if el is comparison_header._p)
    assert tbl_index < comparison_index

    table = document.tables[0]
    assert table.cell(0, 0).text == "Metric"
    assert table.cell(1, 1).text == "1298"
    assert table.cell(2, 1).text == "135"
    assert table.cell(3, 1).text == "97"
    assert table.cell(4, 1).text == "38"
    assert table.cell(5, 1).text == "10 / 1224 = 0.82%"


def test_insert_weakness_retry_audit_table_adds_discrepancy_paragraph():
    document, fig4, table_ii_caption, comparison_header = _results_section_skeleton()

    insert_weakness_retry_audit_table(document, _FAKE_REPORT)

    texts = [p.text for p in document.paragraphs]
    discrepancy = [t for t in texts if "38 dropped pairs" in t]
    assert len(discrepancy) == 1
    assert "26" in discrepancy[0]
    assert "warnings.json" in discrepancy[0]

    body_children = list(document.element.body)
    tbl_index = next(i for i, el in enumerate(body_children) if el.tag.endswith("}tbl"))
    discrepancy_index = next(
        i for i, el in enumerate(body_children) if el is document.paragraphs[texts.index(discrepancy[0])]._p
    )
    comparison_index = next(i for i, el in enumerate(body_children) if el is comparison_header._p)
    assert tbl_index < discrepancy_index < comparison_index


def test_add_evidence_table_captions_labels_both_new_tables():
    document, fig4, table_ii_caption, comparison_header = _results_section_skeleton()
    insert_results_headers_and_ranking_failure_table(document)
    insert_weakness_retry_audit_table(document, _FAKE_REPORT)

    add_evidence_table_captions(document)

    texts = [p.text for p in document.paragraphs]
    table_iv_caption = [t for t in texts if t.startswith("Table IV.")]
    table_v_caption = [t for t in texts if t.startswith("Table V.")]
    assert len(table_iv_caption) == 1
    assert len(table_v_caption) == 1

    body_children = list(document.element.body)
    tbl_indices = [i for i, el in enumerate(body_children) if el.tag.endswith("}tbl")]
    assert len(tbl_indices) == 2
    ranking_failure_tbl_index, weakness_audit_tbl_index = tbl_indices

    def index_of(text: str) -> int:
        p = document.paragraphs[texts.index(text)]
        return body_children.index(p._p)

    table_iv_explanation = [t for t in texts if t.startswith("As shown, all 1,713")]
    assert len(table_iv_explanation) == 1

    # Table IV: [table] -> caption -> explanation, matching the paper's
    # existing Table I/II/III convention (caption + explanation below the table).
    assert ranking_failure_tbl_index < index_of(table_iv_caption[0]) < index_of(table_iv_explanation[0])
    # Table V: [table] -> caption -> (the existing discrepancy-note paragraph
    # already serves as its explanation, added by Task 12 -- not duplicated here).
    assert weakness_audit_tbl_index < index_of(table_v_caption[0])
    discrepancy_index = index_of(
        next(t for t in texts if "38 dropped pairs" in t)
    )
    assert index_of(table_v_caption[0]) < discrepancy_index
