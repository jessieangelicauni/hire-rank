import docx
from docx.shared import Inches
from PIL import Image

from scripts.shorten_paper_20260908 import delete_paragraph, paragraph_element_immediately_before, remove_figure_and_fold_text


def test_delete_paragraph_removes_it_from_the_document():
    document = docx.Document()
    document.add_paragraph("Keep me.")
    target = document.add_paragraph("Delete me.")
    document.add_paragraph("Keep me too.")

    delete_paragraph(target)

    texts = [p.text for p in document.paragraphs]
    assert texts == ["Keep me.", "Keep me too."]


def test_paragraph_element_immediately_before_returns_the_preceding_element():
    document = docx.Document()
    before = document.add_paragraph("Before.")
    anchor = document.add_paragraph("Anchor.")

    element = paragraph_element_immediately_before(document, anchor)

    assert element is before._p


def _tiny_png_path(tmp_path, name="tiny.png"):
    path = tmp_path / name
    Image.new("RGB", (100, 40), color="white").save(path)
    return path


def test_remove_figure_and_fold_text_deletes_image_and_caption(tmp_path):
    document = docx.Document()
    document.add_paragraph("Preceding paragraph.")
    document.add_picture(str(_tiny_png_path(tmp_path)), width=Inches(3.5), height=Inches(2.0))
    document.add_paragraph("Fig. 5. Final Kendall's Tau per job role after convergence.")
    document.add_paragraph("As Fig. 5 shows, values range from 0.930 to 0.979.")
    document.add_paragraph("B. Generation Faithfulness and Contradiction Audit")

    archive_path = tmp_path / "_archive" / "removed.png"

    remove_figure_and_fold_text(
        document,
        caption_marker="Fig. 5. Final Kendall",
        sentence_marker="As Fig. 5 shows",
        new_sentence="Final Kendall's Tau per role ranges from 0.930 to 0.979.",
        archive_path=archive_path,
    )

    texts = [p.text for p in document.paragraphs]
    assert not any(t.startswith("Fig. 5.") for t in texts)
    assert "As Fig. 5 shows" not in "".join(texts)
    assert "Final Kendall's Tau per role ranges from 0.930 to 0.979." in texts
    assert len(document.inline_shapes) == 0
    assert archive_path.exists()


def test_remove_figure_and_fold_text_archives_the_original_image_bytes(tmp_path):
    document = docx.Document()
    png_path = _tiny_png_path(tmp_path)
    document.add_picture(str(png_path), width=Inches(3.5), height=Inches(2.0))
    document.add_paragraph("Fig. 5. Final Kendall's Tau per job role after convergence.")
    document.add_paragraph("As Fig. 5 shows, values range from 0.930 to 0.979.")

    archive_path = tmp_path / "_archive" / "removed.png"

    remove_figure_and_fold_text(
        document,
        caption_marker="Fig. 5. Final Kendall",
        sentence_marker="As Fig. 5 shows",
        new_sentence="Final Kendall's Tau per role ranges from 0.930 to 0.979.",
        archive_path=archive_path,
    )

    assert archive_path.read_bytes() == png_path.read_bytes()
