"""Applies the paper-accuracy corrections from
docs/superpowers/specs/2026-09-05-paper-accuracy-corrections-design.md to
docs/candidate-ranking-paper.docx.

docs/ is not git-tracked, so this script backs up the original file before
overwriting it. Each fix_* function is independently unit-tested against a
small synthetic document in tests/test_apply_paper_corrections.py; main()
runs all of them in order against the real paper.

Run with: uv run python scripts/apply_paper_corrections.py
"""
from __future__ import annotations

import shutil
from pathlib import Path

import docx
from docx.shared import Inches

PAPER_PATH = Path(__file__).resolve().parents[1] / "docs" / "candidate-ranking-paper.docx"
BACKUP_PATH = PAPER_PATH.with_name(PAPER_PATH.stem + ".backup" + PAPER_PATH.suffix)


def find_paragraph(document, contains: str):
    for paragraph in document.paragraphs:
        if contains in paragraph.text:
            return paragraph
    raise ValueError(f"no paragraph found containing: {contains!r}")


def replace_paragraph_text(paragraph, new_text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def insert_paragraph_before(anchor, text: str, bold: bool = False, style: str = "Body Text"):
    try:
        new_paragraph = anchor.insert_paragraph_before(text, style=style)
    except KeyError:
        new_paragraph = anchor.insert_paragraph_before(text)
    if bold and new_paragraph.runs:
        new_paragraph.runs[0].bold = True
    return new_paragraph


def insert_table_before(document, anchor, rows: int, cols: int, col_widths_in: list[float] | None = None):
    table = document.add_table(rows=rows, cols=cols)
    try:
        table.style = document.styles["Table Grid"]
    except KeyError:
        pass
    anchor._p.addprevious(table._tbl)
    if col_widths_in is not None:
        table.autofit = False
        for row in table.rows:
            for cell, width in zip(row.cells, col_widths_in):
                cell.width = Inches(width)
    return table


ABSTRACT_TEXT = (
    "Abstract—Talent-acquisition systems increasingly pair large language models (LLMs) with statistical ranking "
    "to shortlist and order job applicants. The most closely related recent system combines an LLM-driven active "
    "listwise tournament with Plackett-Luce aggregation for human-resources applicant ranking, but it reports no "
    "mechanism against LLM identifier-drift failures during ranking, uses a fixed iteration count that leaves its "
    "convergence unquantified, and offers only prompt-level mitigation against hallucinated assessment claims. "
    "This paper's objective is to validate three pipeline-level mechanisms that close these gaps in an open, "
    "reproducible applicant-ranking system, implemented as a LangGraph-orchestrated pipeline with embedding-based "
    "skill shortlisting. This research's three contributions are: (i) resume-grounded strength/weakness "
    "assessments generated with an automated self-correction retry loop; (ii) ranking of shortlisted applicants "
    "with a position-robust, schema-constrained listwise tournament, adapting the Monte Carlo knowledge-gradient "
    "(MC-KG) subset-selection rule and Bayesian Plackett-Luce aggregation from prior work; and (iii) pool-size-aware "
    "adaptive iteration scaling. The results show that mean Faithfulness reached 0.896 and mean within-repeat "
    "ranking convergence (Kendall-τ) reached 0.953, both measured across all 10 job profiles. In conclusion, "
    "these results — measured with metrics the closest prior system does not report or disclose — offer "
    "transparent, empirical evidence that this research addresses those gaps."
)


def fix_abstract_and_contributions(document) -> None:
    paragraph = find_paragraph(document, "gives no cross-run stability")
    replace_paragraph_text(paragraph, ABSTRACT_TEXT)


def main() -> None:
    shutil.copy2(PAPER_PATH, BACKUP_PATH)
    print(f"Backed up {PAPER_PATH} -> {BACKUP_PATH}")

    document = docx.Document(PAPER_PATH)

    # fix_* calls are added here by later tasks, in this order:
    fix_abstract_and_contributions(document)
    # fix_methods_iii_c_d_e(document)
    # fix_experimental_setup_and_table_ii(document)
    # insert_results_headers_and_ranking_failure_table(document)
    # fix_conclusion_and_limitations(document)
    # insert_weakness_retry_audit_table(document)
    # apply_table_and_figure_layout_fixes(document)

    document.save(PAPER_PATH)
    print(f"Saved corrections to {PAPER_PATH}")


if __name__ == "__main__":
    main()
