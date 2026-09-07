"""Updates docs/candidate-ranking-paper-new's empirical numbers (Kendall-tau,
Faithfulness, weakness-retry-audit) to match run 20260906-150825, and
simplifies the Table IV narrative now that the pipeline logs the retry
breakdown live (see write_retry_audit_report in
candidate_ranking.scoring.assessment) instead of needing an offline replay.

Always rebuilds from docs/candidate-ranking-paper-new.backup (the pristine,
pre-correction file) rather than editing the live file in place, so this
script stays idempotent across re-runs. docs/ is not git-tracked, so that
backup is unrecoverable if lost -- this script never overwrites it.

Also regenerates Fig. 4-6 (docs/figures/table5_ragas_faithfulness.png,
table6_kendall_tau.png, table7_kendall_tau_per_role.png) via
generate_evaluation_figures.py and swaps them into the document, since those
are rendered images, not text the paragraph/table edits below can reach.

Run with: uv run python scripts/update_paper_results_20260906.py
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn

PROJECT_ROOT = Path(__file__).resolve().parents[1]
PAPER_PATH = PROJECT_ROOT / "docs" / "candidate-ranking-paper-new"
BACKUP_PATH = PAPER_PATH.with_name(PAPER_PATH.name + ".backup")
FIGURES_DIR = PROJECT_ROOT / "docs" / "figures"

RUN_ID = "20260906-150825"

# New per-job-profile Faithfulness scores (runs/<RUN_ID>/ragas_faithfulness_report.json,
# per_jd), sorted ascending to match the paper's existing table convention.
NEW_FAITHFULNESS_ROWS = [
    ("full-stack-engineer", "0.764"),
    ("data-scientist-engineer", "0.828"),
    ("ui-ux-designer", "0.831"),
    ("it-security-engineer", "0.872"),
    ("frontend-engineer", "0.894"),
    ("data-engineer", "0.900"),
    ("backend-engineer", "0.917"),
    ("cloud-engineer", "0.917"),
    ("java-developer", "0.933"),
    ("devops-engineer", "0.944"),
]
NEW_FAITHFULNESS_MEAN = "0.880"

# New live weakness-retry-audit numbers (runs/<RUN_ID>/retry_audit.json).
NEW_RETRY_AUDIT = {
    "total_weaknesses_checked": "1268",
    "pairs_with_initial_contradiction": "130",
    "pairs_fixed_by_retry": "88",
    "pairs_dropped": "42",
    "offline_audit_residual": "0 / 1268 = 0.00%",
}


def replace_in_paragraph(paragraph, old: str, new: str) -> bool:
    runs = paragraph.runs
    full_text = "".join(r.text for r in runs)
    if old not in full_text:
        return False
    idx = full_text.index(old)
    end = idx + len(old)
    pos = 0
    inserted = False
    for r in runs:
        r_start = pos
        r_end = pos + len(r.text)
        if r_end > idx and r_start < end:
            local_start = max(0, idx - r_start)
            local_end = min(len(r.text), end - r_start)
            before = r.text[:local_start]
            after = r.text[local_end:]
            if not inserted:
                r.text = before + new + after
                inserted = True
            else:
                r.text = before + after
        pos = r_end
    return True


def replace_everywhere(document, old: str, new: str, expected: int = 1) -> None:
    count = 0
    for paragraph in document.paragraphs:
        if replace_in_paragraph(paragraph, old, new):
            count += 1
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if replace_in_paragraph(paragraph, old, new):
                        count += 1
    if count != expected:
        raise ValueError(f"expected {expected} replacement(s) of {old!r}, got {count}")


def find_paragraph(document, contains: str):
    for paragraph in document.paragraphs:
        if contains in paragraph.text:
            return paragraph
    raise ValueError(f"no paragraph found containing: {contains!r}")


def replace_paragraph_text(paragraph, new_text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def delete_paragraph(paragraph) -> None:
    element = paragraph._p
    element.getparent().remove(element)


def apply_text_and_table_updates(document) -> None:
    replace_everywhere(
        document,
        "mean Faithfulness reached 0.896 and mean within-repeat ranking convergence "
        "(Kendall-τ) reached 0.953",
        "mean Faithfulness reached 0.880 and mean within-repeat ranking convergence "
        "(Kendall-τ) reached 0.957",
    )
    replace_everywhere(
        document,
        "is uniformly high (τ = 0.91–0.98) across all ten job profiles",
        "is uniformly high (τ = 0.93–0.98) across all ten job profiles",
    )
    replace_everywhere(
        document,
        "spans 0.906 (frontend-engineer, n=7) to 0.979 (data-engineer, n=77)",
        "spans 0.930 (frontend-engineer, n=7) to 0.979 (data-engineer, n=77)",
    )
    replace_everywhere(
        document,
        "with 298 of the 300 scored (2 excluded as evaluation failures). The "
        "run-wide mean is 0.896, following successive prompt refinements.",
        "with 300 of the 300 scored (0 excluded as evaluation failures). The "
        "run-wide mean is 0.880, following successive prompt refinements.",
    )
    replace_everywhere(
        document,
        "range from full-stack-engineer (lowest, 0.839) to devops-engineer "
        "(highest, 0.960), bracketing the run-wide mean of 0.896 (Table II).",
        "range from full-stack-engineer (lowest, 0.764) to devops-engineer "
        "(highest, 0.944), bracketing the run-wide mean of 0.880 (Table II).",
    )

    table_ii = document.tables[1]
    assert len(table_ii.rows) == 12, f"unexpected Table II row count: {len(table_ii.rows)}"
    for row_index, (jd, value) in enumerate(NEW_FAITHFULNESS_ROWS, start=1):
        table_ii.cell(row_index, 0).text = jd
        table_ii.cell(row_index, 1).text = value
    table_ii.cell(11, 0).text = "Total Mean"
    table_ii.cell(11, 1).text = NEW_FAITHFULNESS_MEAN

    intro_paragraph = find_paragraph(document, "To recover the before/after-retry breakdown")
    replace_paragraph_text(
        intro_paragraph,
        "The pipeline now records this before/after-retry breakdown directly during "
        "generation, persisting it to the run's own retry_audit.json rather than "
        "requiring a separate offline replay:",
    )

    caption_paragraph = find_paragraph(
        document, "Table IV. Weakness self-correction retry outcomes, from an instrumented replay"
    )
    replace_paragraph_text(
        caption_paragraph,
        "Table IV. Weakness self-correction retry outcomes, recorded live during "
        "generation across all 348 shortlisted pairs.",
    )

    table_iv = document.tables[3]
    assert len(table_iv.rows) == 6, f"unexpected Table IV row count: {len(table_iv.rows)}"
    table_iv.cell(1, 0).text = "Total weaknesses checked"
    table_iv.cell(1, 1).text = NEW_RETRY_AUDIT["total_weaknesses_checked"]
    table_iv.cell(2, 1).text = NEW_RETRY_AUDIT["pairs_with_initial_contradiction"]
    table_iv.cell(3, 1).text = NEW_RETRY_AUDIT["pairs_fixed_by_retry"]
    table_iv.cell(4, 1).text = NEW_RETRY_AUDIT["pairs_dropped"]
    table_iv.cell(5, 0).text = "Residual contradictions found by offline post-hoc audit"
    table_iv.cell(5, 1).text = NEW_RETRY_AUDIT["offline_audit_residual"]

    # The discrepancy paragraph only explained why an offline replay was once
    # needed (and its 38-vs-26 mismatch vs. the original run) -- now that the
    # breakdown is recorded live, that historical justification no longer
    # belongs next to Table IV's live numbers, so it is dropped rather than
    # reworded.
    delete_paragraph(find_paragraph(document, "this replay's 38 dropped pairs"))

    table_v = document.tables[4]
    for row in table_v.rows:
        if row.cells[0].text == "Convergence value disclosure":
            for paragraph in row.cells[2].paragraphs:
                replace_in_paragraph(
                    paragraph,
                    "Numeric Kendall-τ disclosed: 0.91–0.98 across all 10 job profiles",
                    "Numeric Kendall-τ disclosed: 0.93–0.98 across all 10 job profiles",
                )
            break
    else:
        raise ValueError("Table V 'Convergence value disclosure' row not found")

    replace_everywhere(
        document,
        "with a 0.82% residual rate surfaced only by a separate offline post-hoc audit",
        "with a 0.00% residual rate surfaced only by a separate offline post-hoc audit",
    )
    replace_everywhere(
        document,
        "adaptive iteration scaling settles convergence at Kendall-τ 0.91–0.98 "
        "across all ten job profiles",
        "adaptive iteration scaling settles convergence at Kendall-τ 0.93–0.98 "
        "across all ten job profiles",
    )
    replace_everywhere(
        document,
        "Mean Faithfulness reaches 0.896; the residual gap",
        "Mean Faithfulness reaches 0.880; the residual gap",
    )


def swap_evaluation_figures(document) -> None:
    swaps = [
        (3, FIGURES_DIR / "table6_kendall_tau.png", "Fig. 4. Within-repeat convergence"),
        (4, FIGURES_DIR / "table7_kendall_tau_per_role.png", "Fig. 5. Final Kendall"),
        (5, FIGURES_DIR / "table5_ragas_faithfulness.png", "Fig. 6. Faithfulness scores"),
    ]

    shape_idx = 0
    captions_by_shape = {}
    for i, p in enumerate(document.paragraphs):
        if p._p.findall(".//" + qn("w:drawing")):
            caption = ""
            for j in range(i + 1, min(i + 3, len(document.paragraphs))):
                t = document.paragraphs[j].text.strip()
                if t.startswith("Fig."):
                    caption = t
                    break
            captions_by_shape[shape_idx] = caption
            shape_idx += 1

    for idx, new_png, expected_prefix in swaps:
        actual_caption = captions_by_shape.get(idx, "")
        if not actual_caption.startswith(expected_prefix):
            raise ValueError(
                f"inline_shapes[{idx}] caption is {actual_caption!r}, expected prefix "
                f"{expected_prefix!r} -- refusing to overwrite, image order may have changed"
            )
        target_shape = document.inline_shapes[idx]
        rel_id = target_shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        image_part = document.part.related_parts[rel_id]
        image_part._blob = new_png.read_bytes()


def main() -> None:
    if not BACKUP_PATH.exists():
        raise SystemExit(f"missing pristine backup at {BACKUP_PATH}; refusing to run without one")

    subprocess.run(
        [sys.executable, str(PROJECT_ROOT / "scripts" / "generate_evaluation_figures.py")],
        check=True,
        cwd=PROJECT_ROOT,
    )

    document = docx.Document(str(BACKUP_PATH))
    apply_text_and_table_updates(document)
    swap_evaluation_figures(document)
    document.save(str(PAPER_PATH))
    print(f"Wrote {PAPER_PATH}")


if __name__ == "__main__":
    main()
